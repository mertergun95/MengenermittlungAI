"""Mengenermittlung zusammenstellen, als Excel ausgeben und Korrekturen zurücklesen.

Excel-Blätter:
- Mengenermittlung: je OZ ermittelte Menge, LV-Menge, Betrag, Prüfbedarf
- Belege: jede einzelne Leistung mit Quelle, Vorschlag, Konfidenz und Spalte 'OZ korrigiert'
- Prüfliste: Belege ohne Zuordnung, mit niedriger Konfidenz oder mit Hinweisen
Trägt der Prüfer in 'OZ korrigiert' eine OZ (oder 'ok') ein, übernimmt `lerne` sie ins Gedächtnis.
"""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mengen.schema import LVPosition, Mengenzeile
from mengen.zuordnung import FLUECHTIG, Gedaechtnis, Zuordnung, schluessel

GELB = PatternFill("solid", fgColor="FFF2CC")
ROT = PatternFill("solid", fgColor="F8CBAD")
KOPF = PatternFill("solid", fgColor="D9E1F2")

BELEG_SPALTEN = [
    ("Nr", 6), ("OZ Vorschlag", 13), ("OZ korrigiert", 13), ("Konfidenz", 10), ("Methode", 9),
    ("Datum", 11), ("Tätigkeit", 30), ("Material", 28), ("Merkmale", 40), ("Menge", 10), ("Einheit", 7),
    ("Dokument", 30), ("Seite", 6), ("Fundstelle", 60), ("Hinweise", 60), ("Begründung", 50),
    ("Alternativen", 40), ("Schlüssel", 10),
]


def mengenzeilen(zuordnungen: list[Zuordnung]) -> list[Mengenzeile]:
    gruppen: dict[str, list[Zuordnung]] = defaultdict(list)
    for z in zuordnungen:
        if z.oz:
            gruppen[z.oz].append(z)
    zeilen = []
    for oz, zs in sorted(gruppen.items()):
        mengen = [z.eintrag.menge or Decimal(0) for z in zs]
        einheiten = {z.eintrag.einheit for z in zs}
        zeilen.append(Mengenzeile(
            oz=oz, menge=sum(mengen), einheit="/".join(sorted(filter(None, einheiten))),
            rechenweg=f"Summe aus {len(zs)} Belegen" + (
                ": " + " + ".join(f"{m:g}" for m in mengen) if len(zs) <= 12 else ""),
            eintraege=[z.eintrag for z in zs],
            konfidenz=min(z.konfidenz for z in zs),
        ))
    return zeilen


def schreibe_excel(pfad: str | Path, zuordnungen: list[Zuordnung], positionen: list[LVPosition],
                   titel: str = "") -> None:
    wb = openpyxl.Workbook()
    _blatt_mengen(wb.active, zuordnungen, positionen, titel)
    _blatt_belege(wb.create_sheet("Belege"), zuordnungen)
    _blatt_belege(wb.create_sheet("Prüfliste"), [z for z in zuordnungen if z.zu_pruefen], nummern=False,
                  alle=zuordnungen)
    wb.save(pfad)


def _kopfzeile(ws, spalten: list[tuple[str, int]], zeile: int = 1) -> None:
    for i, (name, breite) in enumerate(spalten, start=1):
        c = ws.cell(row=zeile, column=i, value=name)
        c.font, c.fill = Font(bold=True), KOPF
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = ws.cell(row=zeile + 1, column=1)


def _blatt_mengen(ws, zuordnungen: list[Zuordnung], positionen: list[LVPosition], titel: str) -> None:
    ws.title = "Mengenermittlung"
    ws["A1"] = f"Mengenermittlung (Entwurf) {titel}".strip()
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Vorschlag aus Bautagebuch und Verlegeprotokollen. Prüfen und freigeben, bevor abgerechnet wird."
    spalten = [("OZ", 13), ("Kurztext", 45), ("Einheit", 8), ("LV-Menge", 12), ("Ermittelt", 12),
               ("Anteil LV", 10), ("EP", 10), ("Betrag", 14), ("Belege", 8), ("min. Konfidenz", 10),
               ("zu prüfen", 9), ("Rechenweg", 60)]
    _kopfzeile(ws, spalten, zeile=4)
    nach_oz = {m.oz: m for m in mengenzeilen(zuordnungen)}
    pruefen = defaultdict(int)
    for z in zuordnungen:
        if z.oz and z.zu_pruefen:
            pruefen[z.oz] += 1
    r = 5
    for p in positionen:
        m = nach_oz.get(p.oz)
        werte = [p.oz, p.kurztext, p.einheit, p.menge_lv, m.menge if m else None,
                 None, p.ep, None, len(m.eintraege) if m else 0, m.konfidenz if m else None,
                 pruefen[p.oz] or None, m.rechenweg if m else ""]
        for i, w in enumerate(werte, start=1):
            ws.cell(row=r, column=i, value=float(w) if isinstance(w, Decimal) else w)
        ws.cell(row=r, column=6, value=f"=IF(AND(ISNUMBER(D{r}),D{r}<>0,ISNUMBER(E{r})),E{r}/D{r},\"\")")
        ws.cell(row=r, column=8, value=f"=IF(AND(ISNUMBER(E{r}),ISNUMBER(G{r})),E{r}*G{r},\"\")")
        for spalte, fmt in ((4, "#,##0.00"), (5, "#,##0.00"), (6, "0%"), (7, "#,##0.00"), (8, "#,##0.00 €"),
                            (10, "0%")):
            ws.cell(row=r, column=spalte).number_format = fmt
        if pruefen[p.oz]:
            ws.cell(row=r, column=11).fill = GELB
        r += 1
    ws.cell(row=r, column=7, value="Summe").font = Font(bold=True)
    ws.cell(row=r, column=8, value=f"=SUM(H5:H{r - 1})").number_format = "#,##0.00 €"
    ws.cell(row=r, column=8).font = Font(bold=True)
    ohne = [z for z in zuordnungen if not z.oz]
    if ohne:
        ws.cell(row=r + 2, column=1, value=f"{len(ohne)} Belege ohne passende LV-Position, siehe Prüfliste "
                                            "(mögliche Nachtragsleistungen).").font = Font(italic=True)
    # Mehr als 110 % der LV-Menge: Nachtrag oder Fehlzuordnung prüfen
    ws.conditional_formatting.add(f"F5:F{r - 1}", CellIsRule(operator="greaterThan", formula=["1.1"], fill=ROT))
    ws.auto_filter.ref = f"A4:{get_column_letter(len(spalten))}{r - 1}"


def _blatt_belege(ws, zuordnungen: list[Zuordnung], nummern: bool = True, alle: list[Zuordnung] | None = None) -> None:
    _kopfzeile(ws, BELEG_SPALTEN)
    index = {id(z): i for i, z in enumerate(alle or zuordnungen, start=1)}
    for r, z in enumerate(zuordnungen, start=2):
        e = z.eintrag
        werte = [
            index[id(z)], z.oz, None, z.konfidenz, z.methode, e.datum, e.taetigkeit, e.material,
            ", ".join(f"{k}={v}" for k, v in e.merkmale.items() if k not in FLUECHTIG or k == "bezeichnung"),
            float(e.menge) if e.menge is not None else None, e.einheit, e.quelle.dokument, e.quelle.seite,
            (e.quelle.zitat or "").replace("\n", " | "), "; ".join(e.hinweise), z.begruendung,
            "; ".join(f"{k.oz} ({k.score:.2f})" for k in z.kandidaten if k.oz != z.oz)[:200], schluessel(e),
        ]
        for i, w in enumerate(werte, start=1):
            c = ws.cell(row=r, column=i, value=w)
            c.alignment = Alignment(vertical="top", wrap_text=i in (14, 15, 16))
        ws.cell(row=r, column=4).number_format = "0%"
        ws.cell(row=r, column=6).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=3).fill = GELB
        if not z.oz:
            ws.cell(row=r, column=2).fill = ROT
        elif z.konfidenz < 0.6:
            ws.cell(row=r, column=4).fill = GELB
    ws.column_dimensions[get_column_letter(len(BELEG_SPALTEN))].hidden = True
    ws.auto_filter.ref = f"A1:{get_column_letter(len(BELEG_SPALTEN))}{max(2, len(zuordnungen) + 1)}"


def lerne(pfad: str | Path, gedaechtnis: Gedaechtnis) -> int:
    """'OZ korrigiert' aus dem Blatt Belege übernehmen: OZ eintragen oder 'ok' für den Vorschlag."""
    ws = openpyxl.load_workbook(pfad, data_only=True)["Belege"]
    kopf = [c.value for c in ws[1]]
    i_vorschlag, i_korr, i_schl = (kopf.index(n) for n in ("OZ Vorschlag", "OZ korrigiert", "Schlüssel"))
    anzahl = 0
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        korr = str(zeile[i_korr] or "").strip()
        if not korr or not zeile[i_schl]:
            continue
        oz = zeile[i_vorschlag] if korr.lower() == "ok" else korr
        if oz:
            gedaechtnis.merke(zeile[i_schl], str(oz))
            anzahl += 1
    gedaechtnis.speichere()
    return anzahl
