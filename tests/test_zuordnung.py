from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from mengen.aufmass import lerne, mengenzeilen, schreibe_excel
from mengen.gaeb import lese_lv, verbinde, vereinige
from mengen.schema import Leistungseintrag, LVPosition, Quelle
from mengen.zuordnung import Gedaechtnis, Zuordner, kennwerte, pruefe

FIX = Path(__file__).parent / "fixtures"
D = Decimal


@pytest.fixture(scope="module")
def positionen():
    return vereinige(
        verbinde(lese_lv(FIX / "Basel_LV.x83"), lese_lv(FIX / "Basel_LV.x84")),
        verbinde(lese_lv(FIX / "Basel_Nachtrag_LV.x83"), lese_lv(FIX / "Basel_Nachtrag_LV.x84")),
    )


def leistung(taetigkeit, menge, einheit, **merkmale):
    return Leistungseintrag(taetigkeit=taetigkeit, menge=D(menge), einheit=einheit,
                            merkmale={k: str(v) for k, v in merkmale.items()},
                            quelle=Quelle(dokument="test.pdf", seite=1))


@pytest.mark.parametrize("kurztext, langtext, erwartet", [
    ("Kabel einziehen, D über 25-40 mm", "", {"durchmesser_mm": (D(25), D(40), True)}),
    ("Kabel einziehen, D bis 25 mm", "", {"durchmesser_mm": (None, D(25), False)}),
    ("Kabel einziehen", "Durchmesser über 40 mm", {"durchmesser_mm": (D(40), None, True)}),
    ("Kabeltrog Größe IV-V a.D. öffnen", "", {"groesse": (D(4), D(5), False)}),
    ("Kabeschacht Gr V öffnen", "", {"groesse": (D(5), D(5), False)}),
    ("Schachtanschlussbausatz öffnen", "", {}),
])
def test_kennwerte_aus_lv_text(kurztext, langtext, erwartet):
    assert kennwerte(LVPosition(oz="1", kurztext=kurztext, langtext=langtext, einheit="m")) == erwartet


def test_pruefe_bereich():
    ueber_25_bis_40 = (D(25), D(40), True)
    assert pruefe(ueber_25_bis_40, (D(39), D(39))) == "ja"
    assert pruefe(ueber_25_bis_40, (D(25), D(25))) == "nein"   # 'über 25' schließt 25 aus
    assert pruefe(ueber_25_bis_40, (D(41), D(41))) == "nein"
    assert pruefe(ueber_25_bis_40, (D(38), D(42))) == "teils"


def test_kabel_nach_durchmesser(positionen):
    z = Zuordner(positionen)
    assert z.ordne_zu(leistung("Kabel einziehen", 175, "m", durchmesser_mm="12.0")).oz == "21.03.0010"
    assert z.ordne_zu(leistung("Kabel einziehen", 175, "m", durchmesser_mm="39.0")).oz == "21.02.0010"
    # 41 mm passt in keine Position -> kein Vorschlag statt Raten
    zu_gross = z.ordne_zu(leistung("Kabel einziehen", 175, "m", durchmesser_mm="41.0"))
    assert zu_gross.oz is None and zu_gross.zu_pruefen
    # Bereich überlappt die Grenze -> Vorschlag, aber zur Prüfung
    unklar = z.ordne_zu(leistung("Kabel einziehen", 175, "m", durchmesser_mm="38.0–42.0"))
    assert unklar.oz == "21.02.0010" and unklar.konfidenz <= 0.4


def test_kanal_groesse_ueber_synonyme(positionen):
    z = Zuordner(positionen)
    r = z.ordne_zu(leistung("Kabelkanal öffnen und schließen", 144, "m", art="BKK", groesse=2, lage="innenliegend"))
    assert r.oz == "99.01.0020" and r.konfidenz > 0.9
    # Größe 4 passt auf 'IV' und 'IV-V' -> unsicher
    r4 = z.ordne_zu(leistung("Kabelkanal öffnen und schließen", 50, "m", art="BKK", groesse=4))
    assert r4.oz == "99.01.0040" and r4.konfidenz < 0.6


def test_nicht_pruefbarer_kennwert_senkt_konfidenz(positionen):
    r = Zuordner(positionen).ordne_zu(leistung("Schacht öffnen und schließen", 3, "St", groesse="klein"))
    assert r.oz == "99.01.0060" and r.konfidenz <= 0.4


def test_einheit_passt_nicht(positionen):
    r = Zuordner(positionen).ordne_zu(leistung("Schacht öffnen und schließen", 30, "m"))
    assert all("Einheit m ≠ St" in " ".join(k.gruende) for k in r.kandidaten if k.oz == "99.01.0060")


class FakeLLM:
    def __init__(self, antwort):
        self.antwort, self.prompts = antwort, []

    def json(self, prompt):
        self.prompts.append(prompt)
        return self.antwort


def test_llm_nur_bei_unsicherheit_und_nur_kandidaten(positionen):
    llm = FakeLLM({"oz": "99.01.0050", "begruendung": "IV-V passt besser"})
    z = Zuordner(positionen, llm=llm)
    sicher = z.ordne_zu(leistung("Kabelkanal öffnen und schließen", 10, "m", art="BKK", groesse=1))
    assert sicher.methode == "regel" and not llm.prompts
    unsicher = z.ordne_zu(leistung("Kabelkanal öffnen und schließen", 10, "m", art="BKK", groesse=4))
    assert (unsicher.oz, unsicher.methode) == ("99.01.0050", "llm") and "99.01.0040" in llm.prompts[0]
    erfunden = Zuordner(positionen, llm=FakeLLM({"oz": "77.77.7777"})).ordne_zu(
        leistung("Kabelkanal öffnen und schließen", 10, "m", art="BKK", groesse=4))
    assert erfunden.oz == "99.01.0040"  # OZ außerhalb der Kandidaten wird ignoriert


def test_excel_und_lernen(positionen, tmp_path):
    eintraege = [
        leistung("Kabel einziehen", 175, "m", durchmesser_mm="39.0", bezeichnung="S1651"),
        leistung("Kabel einziehen", 210, "m", durchmesser_mm="27.0", bezeichnung="S2505"),
        leistung("Schacht öffnen und schließen", 3, "St", groesse="klein"),
        leistung("Kabel umverlegen", 374, "m"),
    ]
    gedaechtnis = Gedaechtnis(tmp_path / "projekt.json")
    z = Zuordner(positionen, gedaechtnis)
    zuordnungen = [z.ordne_zu(e) for e in eintraege]
    assert [(m.oz, m.menge) for m in mengenzeilen(zuordnungen)] == [("21.02.0010", D(385)), ("99.01.0060", D(3))]

    pfad = tmp_path / "aufmass.xlsx"
    schreibe_excel(pfad, zuordnungen, positionen)
    wb = openpyxl.load_workbook(pfad)
    assert wb.sheetnames == ["Mengenermittlung", "Belege", "Prüfliste"]
    kopf = [c.value for c in wb["Belege"][1]]
    assert wb["Belege"].max_row == 5
    pruef_oz = [r[1] for r in wb["Prüfliste"].iter_rows(min_row=2, values_only=True)]
    assert "99.01.0060" in pruef_oz and None in pruef_oz  # Schacht unsicher, Umverlegung ohne Position

    # Prüfer bestätigt den Schacht-Vorschlag ('ok') und ordnet die Umverlegung einer OZ zu
    ws = wb["Belege"]
    korr = kopf.index("OZ korrigiert") + 1
    ws.cell(row=4, column=korr, value="ok")
    ws.cell(row=5, column=korr, value="99.01.0070")
    wb.save(pfad)
    assert lerne(pfad, gedaechtnis) == 2

    neu = Zuordner(positionen, Gedaechtnis(tmp_path / "projekt.json"))
    schacht = neu.ordne_zu(leistung("Schacht öffnen und schließen", 5, "St", groesse="klein"))
    assert (schacht.oz, schacht.methode, schacht.zu_pruefen) == ("99.01.0060", "gelernt", False)
    assert neu.ordne_zu(leistung("Kabel umverlegen", 100, "m")).oz == "99.01.0070"
